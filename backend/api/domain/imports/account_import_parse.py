"""Bounded, offline account document parsing; never reconciles or writes facts."""

from __future__ import annotations

import csv
import re
from datetime import date, datetime, time
from decimal import Decimal, InvalidOperation
from io import BytesIO, StringIO
from pathlib import PurePath
from typing import Any
from xml.etree import ElementTree as ET
from zipfile import BadZipFile, ZipFile
from domain.imports import account_import_ths as ths

from openpyxl import load_workbook
from openpyxl.utils.cell import column_index_from_string

MAX_ROWS, MAX_SHEETS, MAX_CELLS = 2000, 10, 100000
MAX_INPUT_BYTES, MAX_EXPANDED_BYTES, MAX_ZIP_RATIO = 10 * 1024**2, 40 * 1024**2, 100
HEADER_SCAN = 20
ALIASES = {
    "business_date": "日期 业务日期 成交日期 交易日期 持仓日期 资产日期",
    "trade_time": "成交时间 交易时间", "confirmed_date": "确认日期 确认日",
    "kind": "记录类型", "asset_type": "资产类型 证券类型 产品类型",
    "symbol": "证券代码 股票代码 基金代码 标的代码", "name": "证券名称 股票名称 基金名称 标的名称",
    "currency": "币种 货币", "action": "方向 买卖方向 操作 业务类型 交易类型 业务名称",
    "status": "状态 交易状态 确认状态", "quantity": "持仓数量 份额 总份额 成交数量 成交份额 数量",
    "price": "成交价格 成交价 价格 单价", "gross_amount": "成交金额 发生金额 交易金额",
    "fee_amount": "手续费 费用 税费 税费合计 金额税费total", "tax_amount": "印花税 税额", "fee_scope": "费用口径",
    "net_cash": "净收付 净发生额 现金发生额",
    "market_value": "市值 持仓市值 资产市值", "cost": "成本金额 持仓成本 报表成本 摊薄成本金额",
    "cost_basis_type": "成本类型 成本口径", "counter_account_code": "对方账户 对方账号",
    "source_ref": "成交编号 成交序号 流水号 交易流水号 申请编号 订单号",
    "asset_id": "标的ID 资产ID", "note": "备注 说明", "duplicate_decision": "重复处理",
}
FIELDS = tuple(ALIASES)
HEADER_FIELDS = {alias: field for field, aliases in ALIASES.items() for alias in [field, *aliases.split()]}
NUMBERS = {"quantity", "price", "gross_amount", "fee_amount", "tax_amount", "net_cash", "market_value", "cost"}
AGGREGATE_FEES = {"税费", "税费合计", "金额税费total"}
ENUMS = {
    "kind": {"trade": "交易", "holding": "持仓", "cash": "现金"},
    "asset_type": {"stock": "股票 A股", "fund": "基金 ETF", "cash": "现金"},
    "action": {"buy": "买入 证券买入", "sell": "卖出 证券卖出", "subscribe": "申购 认购 基金申购 基金买入",
               "redeem": "赎回 基金赎回 基金卖出", "dividend": "分红 红利 现金分红", "fee": "费用 手续费",
               "deposit": "入金 存入 银证转入", "withdraw": "出金 取出 银证转出", "balance": "余额",
               "transfer": "转账 转入 转出 划转", "income": "收入 收入到账", "expense": "支出 消费付款"},
    "status": {"confirmed": "已确认 确认成功 成交 已成交 全部成交 交易成功", "pending": "待确认 确认中 处理中 待成交"},
    "cost_basis_type": {"holding_cost": "持仓成本", "diluted_cost": "摊薄成本 摊薄成本价", "unknown": "未知 报表成本"},
    "currency": {"CNY": "人民币 RMB ￥ ¥"},
    "fee_scope": {"total_charges": "汇总税费", "separate": "分列"},
}
MISSING = {"", "-", "--", "N/A", "null", "None", "未知"}


def _text(value: Any) -> str:
    if isinstance(value, (date, time)):
        return value.isoformat()
    return str(value).strip() if value is not None else ""


def _issue(field: str | None, code: str, message: str, severity: str = "error") -> dict:
    return {"field": field, "code": code, "message": message, "severity": severity}


def _value(field: str, value: Any) -> str | None:
    raw = _text(value)
    if raw in MISSING and not (field == "cost_basis_type" and raw == "未知"):
        return None
    if field in NUMBERS:
        currency = r"^(?:CNY|RMB|人民币|[￥¥])\s*|\s*元$"
        raw = re.sub(currency, "", raw, flags=re.IGNORECASE).strip()
        if raw.startswith("(") and raw.endswith(")"):
            raw = "-" + re.sub(currency, "", raw[1:-1].strip(), flags=re.IGNORECASE).strip()
        if len(raw) > 128 or not re.fullmatch(r"[+-]?(?:\d{1,3}(?:,\d{3})+|\d+|\d*\.\d+)(?:\.\d+)?(?:[eE][+-]?\d{1,3})?", raw):
            raise ValueError("Invalid decimal")
        number = Decimal(raw.replace(",", ""))
        exponent = number.as_tuple().exponent
        if not number.is_finite() or abs(number.adjusted()) > 100 or not isinstance(exponent, int) or exponent < -100:
            raise ValueError("Non-finite or oversized decimal")
        return format(number, "f")
    if field in {"business_date", "confirmed_date"}:
        parts = re.split(r"[T\s]+", raw, maxsplit=1)
        if len(parts) > 1:
            time.fromisoformat(parts[1])
        raw = parts[0].replace("年", "-").replace("月", "-").replace("日", "")
        for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d", "%Y%m%d"):
            try:
                return datetime.strptime(raw, fmt).date().isoformat()
            except ValueError:
                pass
        raise ValueError("Invalid date")
    if field == "trade_time":
        parsed = time.fromisoformat(re.split(r"[T ]", raw)[-1])
        if parsed.tzinfo is not None:
            raise ValueError("Trade time must be a local wall time")
        return parsed.replace(microsecond=0).isoformat()
    if field == "symbol":
        if not re.fullmatch(r"[0-9]{1,6}", raw):
            raise ValueError("Expected a mainland six-digit code")
        return raw.zfill(6)
    if field in ENUMS:
        for canonical, aliases in ENUMS[field].items():
            if raw.lower() == canonical.lower() or raw in aliases.split():
                return canonical
        raise ValueError("Unmapped value")
    return raw


def _xlsx(content: bytes) -> list[tuple[str, list[list[str]]]]:
    # Inspect the bounded archive before openpyxl can expand sparse sheets or links.
    try:
        with ZipFile(BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > 1000 or sum(e.file_size for e in entries) > MAX_EXPANDED_BYTES:
                raise ValueError("XLSX expanded size limit exceeded")
            if len({e.filename for e in entries}) != len(entries):
                raise ValueError("Duplicate ZIP members")
            cells = sheets = 0
            for entry in entries:
                name = entry.filename.lower()
                if entry.flag_bits & 1 or entry.file_size > max(entry.compress_size, 1) * MAX_ZIP_RATIO:
                    raise ValueError("Encrypted or excessive ZIP compression ratio")
                if "vbaproject" in name or "macrosheet" in name or "externallink" in name or name.endswith(".bin"):
                    raise ValueError("Macros and external links are unsupported")
                if not name.endswith((".xml", ".rels")):
                    continue
                xml = archive.read(entry)
                decoded = xml.decode("utf-8-sig")
                if "\x00" in decoded or any(token in decoded.upper() for token in ("<!DOCTYPE", "<!ENTITY")) or "macroenabled" in decoded.lower():
                    raise ValueError("Unsafe XLSX XML or macro content")
                worksheet = name.startswith("xl/worksheets/") and name.endswith(".xml")
                sheets += int(worksheet)
                if sheets > MAX_SHEETS:
                    raise ValueError("Sheet limit exceeded")
                for _, node in ET.iterparse(BytesIO(xml), events=("end",)):
                    tag = node.tag.rsplit("}", 1)[-1]
                    if tag in {"f", "calculatedColumnFormula", "totalsRowFormula", "externalReferences"} or node.get("TargetMode", "").lower() == "external":
                        raise ValueError("Formulas and external links are unsupported")
                    if worksheet and tag == "c":
                        cells += 1
                        if cells > MAX_CELLS:
                            raise ValueError("Cell limit exceeded")
                        ref = node.get("r")
                        match = re.fullmatch(r"([A-Z]{1,3})([1-9]\d{0,6})", ref) if ref else None
                        if ref and (not match or column_index_from_string(match[1]) > 16384 or int(match[2]) > MAX_ROWS + HEADER_SCAN):
                            raise ValueError("Invalid or oversized cell coordinate")
                    if worksheet and tag == "row" and int(node.get("r", "0")) > MAX_ROWS + HEADER_SCAN:
                        raise ValueError("Worksheet row limit exceeded")
                    node.clear()
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
        try:
            if len(workbook.sheetnames) > MAX_SHEETS:
                raise ValueError("Sheet limit exceeded")
            tables: list[tuple[str, list[list[str]]]] = []
            scanned = 0
            for sheet in workbook.worksheets:
                sheet.reset_dimensions()
                rows = []
                for index, row in enumerate(sheet.iter_rows(), 1):
                    scanned += len(row)
                    if index > MAX_ROWS + HEADER_SCAN or scanned > MAX_CELLS:
                        raise ValueError("Worksheet row/cell limit exceeded")
                    if any(cell.data_type in {"f", "e"} for cell in row):
                        raise ValueError("Formula or error cell is unsupported")
                    rows.append([_text(cell.value) for cell in row])
                tables.append((sheet.title, rows))
            return tables
        finally:
            workbook.close()
    except (BadZipFile, ET.ParseError, KeyError, OSError, TypeError, UnicodeError) as exc:
        raise ValueError("Malformed XLSX document") from exc


def _csv(content: bytes) -> list[tuple[None, list[list[str]]]]:
    try:
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = content.decode("gb18030")
        if "\x00" in text or text.lstrip().startswith(("{", "[")):
            raise ValueError("Binary/JSON payload is not a tabular document")
        sample = text[:8192]
        try:
            delimiter = csv.Sniffer().sniff(sample, delimiters=",\t;|").delimiter
        except csv.Error:
            delimiter = max(",\t;|", key=lambda char: max((line.count(char) for line in sample.splitlines()[:HEADER_SCAN]), default=0))
        rows: list[list[str]] = []
        cells = 0
        for row in csv.reader(StringIO(text, newline=""), delimiter=delimiter, strict=True):
            cells += len(row)
            if len(rows) >= MAX_ROWS + HEADER_SCAN or cells > MAX_CELLS:
                raise ValueError("CSV row/cell limit exceeded")
            if any(cell.lstrip().startswith("=") for cell in row):
                raise ValueError("Formula cells are unsupported")
            rows.append(row)
        return [(None, rows)]
    except (UnicodeError, csv.Error) as exc:
        raise ValueError("Malformed CSV/TSV document") from exc


def parse_document(content: bytes, filename: str, profile: str, kind: str, mapping: dict | None = None) -> dict:
    """Return rows/headers/warnings/suggested_mapping without I/O or input mutation.

    mapping is canonical field -> exact source header (no constants or indexes).
    source.row is one-based; raw_text reconstructs all source cells with tabs.
    Warnings and row issues share {field, code, message, severity}. Unsafe files,
    invalid arguments and invalid mapping shapes raise ValueError; missing mapped
    headers instead flag headers_changed and leave that field unavailable.
    """
    if profile not in {"tonghuashun", "tencent_licaitong"} or kind not in {"trade", "holding", "cash"}:
        raise ValueError("Unsupported profile or kind")
    if not isinstance(content, bytes) or not content or len(content) > MAX_INPUT_BYTES:
        raise ValueError("Empty, invalid or oversized input")
    if mapping is not None and (not isinstance(mapping, dict) or any(
        field not in FIELDS or not isinstance(header, str) or not header.strip() for field, header in mapping.items()
    )):
        raise ValueError("Mapping must contain canonical fields and nonempty source headers")
    mapping = dict(mapping or {})
    if len(set(mapping.values())) != len(mapping):
        raise ValueError("A source header cannot map to multiple fields")
    result: dict[str, Any] = {"rows": [], "headers": [], "warnings": [], "suggested_mapping": {}}
    suffix = PurePath(filename).suffix.lower()
    if suffix in {".png", ".jpg", ".jpeg"}:
        result["warnings"].append(_issue(None, "ocr_pending", "Image parsing awaits OCR; no values inferred", "warning"))
        return result
    if suffix not in {".csv", ".tsv", ".xlsx", ".xls"}:
        raise ValueError("Unsupported file type")
    if suffix == ".xls" and (profile != "tonghuashun" or kind != "trade" or content.lstrip().startswith((b"\xd0\xcf\x11\xe0", b"PK", b"<"))):
        raise ValueError("仅支持同花顺原始文本型.xls交易报表；二进制Excel或其他资料请使用XLSX")
    first_headers = None
    for sheet, table in (_xlsx(content) if suffix == ".xlsx" else _csv(content)):
        if not any(any(_text(v) for v in row) for row in table):
            continue
        score = lambda row: sum(_text(v) in HEADER_FIELDS or _text(v) in mapping.values() for v in row)
        header_index = max(range(min(HEADER_SCAN, len(table))), key=lambda i: score(table[i]))
        headers = [_text(v) for v in table[header_index]]
        native = profile == "tonghuashun" and kind == "trade" and ths.matches(headers)
        if suffix == ".xls" and not native:
            raise ValueError("未识别为同花顺交易报表，不按.xls后缀信任文件内容")
        effective_mapping = {**{k: v for k, v in ths.MAPPING.items() if v in headers}, **mapping} if native else mapping
        if len([h for h in headers if h]) != len({h for h in headers if h}):
            raise ValueError("Duplicate column headers")
        if not score(headers):
            raise ValueError("No recognizable column headers")
        changed = first_headers is not None and headers != first_headers
        first_headers = headers if first_headers is None else first_headers
        unknown = [h for h in headers if h and h not in HEADER_FIELDS and h not in effective_mapping.values() and not (native and h in ths.EXTRA_HEADERS)]
        if changed or unknown or any(h not in headers for h in mapping.values()):
            result["warnings"].append(_issue(None, "headers_changed", f"Review headers in {sheet or filename}: {headers}", "warning"))
        result["headers"].extend(h for h in headers if h and h not in result["headers"])
        declared = set(effective_mapping) | {HEADER_FIELDS[h] for h in headers if h in HEADER_FIELDS}
        columns = {}
        for field in FIELDS:
            matches = [i for i, h in enumerate(headers) if h and (h == effective_mapping[field] if field in effective_mapping else HEADER_FIELDS.get(h) == field and h not in effective_mapping.values())]
            if len(matches) == 1:
                columns[field] = matches[0]
                result["suggested_mapping"].setdefault(field, headers[matches[0]])
            elif len(matches) > 1:
                result["warnings"].append(_issue(field, "ambiguous_header", "Choose one source header explicitly", "warning"))
        for index, row in enumerate(table[header_index + 1:], header_index + 2):
            if not any(_text(v) for v in row):
                continue
            if len(result["rows"]) >= MAX_ROWS:
                raise ValueError("Data row limit exceeded")
            fields, issues = dict.fromkeys(FIELDS), []
            if any(_text(v) for v in row[len(headers):]) or any(_text(v) and not headers[i] for i, v in enumerate(row[:len(headers)])):
                issues.append(_issue(None, "row_shape", "Values have no corresponding column header"))
            for field, column in columns.items():
                raw = row[column] if column < len(row) else None
                try:
                    fields[field] = _value(field, raw)
                except (ValueError, InvalidOperation):
                    issues.append(_issue(field, "invalid_value", f"Unrecognized {field}: {_text(raw)}"))
            if fields["kind"] not in {None, kind}:
                issues.append(_issue("kind", "kind_mismatch", "Source kind differs from requested kind"))
            if "kind" not in declared:
                fields["kind"] = kind
            if "currency" not in declared:
                fields["currency"] = "CNY"
            fee_header = headers[columns["fee_amount"]] if "fee_amount" in columns else None
            if fee_header in AGGREGATE_FEES:
                if fields["fee_scope"] not in {None, "total_charges"}:
                    issues.append(_issue("fee_scope", "scope_conflict", "Aggregate header conflicts with source fee scope"))
                fields["fee_scope"] = "total_charges"
            elif fee_header == "手续费" and "fee_scope" not in declared:
                fields["fee_scope"] = "separate"
            if fields["fee_scope"] == "total_charges":
                if fields["tax_amount"] is not None:
                    issues.append(_issue("tax_amount", "aggregate_breakdown", "Aggregate charges do not establish a tax breakdown", "warning"))
                fields["tax_amount"] = None
            if "asset_type" not in declared:
                identifiers = {headers[columns[f]] for f in ("symbol", "name") if f in columns}
                fund = bool(identifiers & {"基金代码", "基金名称"})
                stock = bool(identifiers & {"股票代码", "股票名称"})
                if kind == "cash":
                    fields["asset_type"] = "cash"
                elif fund != stock:
                    fields["asset_type"] = "fund" if fund else "stock"
            if profile == "tencent_licaitong" and fields["asset_type"] == "fund" and fields["action"] in {"buy", "sell"}:
                fields["action"] = {"buy": "subscribe", "sell": "redeem"}[fields["action"]]
            if fields["asset_type"] == "stock" and fields["action"] in {"subscribe", "redeem"}:
                fields["action"] = None
                issues.append(_issue("action", "ambiguous_action", "Stock subscription/redemption requires review"))
            executed = "成交日期" in headers and any(h in headers for h in ("成交数量", "成交金额", "成交价格"))
            if "status" not in declared and executed and fields["business_date"] and any(fields[f] is not None for f in ("quantity", "price", "gross_amount")) and fields["action"] in {"buy", "sell", "subscribe", "redeem"}:
                fields["status"] = "confirmed"
            cost_header = headers[columns["cost"]] if "cost" in columns else None
            diluted = any("摊薄成本" in h for h in headers)
            if "cost_basis_type" not in declared and (fields["cost"] is not None or diluted):
                fields["cost_basis_type"] = "holding_cost" if cost_header == "持仓成本" else "diluted_cost" if cost_header == "摊薄成本金额" or (diluted and cost_header is None) else "unknown"
            if diluted or fields["cost_basis_type"] == "diluted_cost":
                issues.append(_issue("cost_basis_type", "diluted_cost", "Diluted cost is not holding cost", "warning"))
            if fields["cost_basis_type"] == "unknown":
                issues.append(_issue("cost_basis_type", "unknown_cost_basis", "Report cost basis is not established", "warning"))
            source = {"sheet": sheet, "row": index, "raw_text": "\t".join(row)}
            if native:
                raw_row = {h: row[i] if i < len(row) else None for i, h in enumerate(headers) if h}
                native_source, native_issues = ths.normalize(fields, raw_row, _value, mapping)
                source.update(native_source)
                issues = [i for i in issues if not (i["field"] == "action" and fields.get("action") is not None)]
                issues.extend(native_issues)
            required = ["business_date"] + (["action", "status"] if kind == "trade" else ["quantity"] if kind == "holding" else ["action"])
            for field in required:
                if fields[field] is None:
                    issues.append(_issue(field, "missing", f"Missing {field}"))
            if kind != "cash" and not (fields["symbol"] or fields["name"]):
                issues.append(_issue("symbol", "missing", "Missing both symbol and name"))
            if kind != "cash" and fields["asset_type"] is None and not native:
                issues.append(_issue("asset_type", "missing", "Asset type needs explicit evidence"))
            result["rows"].append({"fields": fields, "source": source, "issues": issues})
    return result
