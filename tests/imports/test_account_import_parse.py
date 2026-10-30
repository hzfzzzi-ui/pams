from __future__ import annotations

import csv
import sys
from copy import deepcopy
from datetime import datetime, time
from io import BytesIO, StringIO
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from openpyxl import Workbook

API_ROOT = Path(__file__).resolve().parents[2] / "backend" / "api"
if str(API_ROOT) not in sys.path:
    sys.path.insert(0, str(API_ROOT))

from domain.imports import account_import_parse as parser


def csv_bytes(rows, *, encoding="utf-8-sig", delimiter=","):
    stream = StringIO(newline="")
    csv.writer(stream, delimiter=delimiter).writerows(rows)
    return stream.getvalue().encode(encoding)


def xlsx_bytes(rows, *, sheets=1):
    book = Workbook()
    for index in range(sheets):
        sheet = book.active if index == 0 else book.create_sheet()
        sheet.title = f"Sheet{index + 1}"
        for row in rows:
            sheet.append(row)
    content = BytesIO()
    book.save(content)
    book.close()
    return content.getvalue()


def parse(rows, *, kind="trade", profile="tonghuashun", mapping=None):
    return parser.parse_document(csv_bytes(rows), "account.csv", profile, kind, mapping)


def codes(result):
    return {(issue["field"], issue["code"]) for row in result["rows"] for issue in row["issues"]}


def zip_bytes(entries):
    stream = BytesIO()
    with ZipFile(stream, "w", compression=ZIP_DEFLATED) as archive:
        for name, content in entries:
            archive.writestr(name, content)
    return stream.getvalue()


def test_execution_contract_leading_zero_distinct_refs_and_no_computations():
    result = parse([
        ["成交日期", "成交时间", "证券代码", "证券名称", "证券类型", "买卖方向", "成交数量", "成交价格", "成交金额", "手续费", "印花税", "净收付", "成交编号"],
        ["20261028", "09:30:01", "000001", "测试股票", "股票", "买入", "100", "10", "999", "", "0", "(1,002.50)", "00001"],
        ["2026/10/28", "09:31", "000001", "测试股票", "股票", "买入", "100", "10", "999", "0", "", "-1002.50", "00002"],
    ])
    assert set(result) == {"rows", "headers", "warnings", "suggested_mapping"}
    one, two = (row["fields"] for row in result["rows"])
    assert set(one) == set(parser.FIELDS)
    assert all(v is None or isinstance(v, str) for v in one.values())
    assert one["symbol"] == two["symbol"] == "000001"
    assert [one["source_ref"], two["source_ref"]] == ["00001", "00002"]
    assert one["business_date"] == "2026-10-28" and one["trade_time"] == "09:30:01"
    assert two["trade_time"] == "09:31:00"
    assert one["status"] == "confirmed" and one["currency"] == "CNY" and one["kind"] == "trade"
    assert one["confirmed_date"] is None and one["gross_amount"] == "999"  # Not quantity * price.
    assert one["fee_amount"] is None and two["fee_amount"] == "0"
    assert one["tax_amount"] == "0" and two["tax_amount"] is None
    assert one["net_cash"] == two["net_cash"] == "-1002.50"
    assert result["rows"][0]["source"]["row"] == 2 and result["rows"][0]["source"]["sheet"] is None
    assert "00001" in result["rows"][0]["source"]["raw_text"]
    assert not result["rows"][0]["issues"]


@pytest.mark.parametrize("encoding,delimiter", [("utf-8-sig", ","), ("gb18030", ","), ("utf-8", "\t"), ("utf-8", ";")])
@pytest.mark.parametrize("direction,expected", [("买入", "subscribe"), ("卖出", "redeem"), ("分红", "dividend")])
def test_fund_profiles_preamble_encodings_delimiters(encoding, delimiter, direction, expected):
    content = csv_bytes([
        ["腾讯理财通交易记录"], [],
        ["日期", "确认日期", "基金代码", "基金名称", "交易类型", "状态", "份额", "发生金额"],
        ["2026年10月27日", "2026.10.28", "000002", "测试基金", direction, "已确认", "12.50", "CNY 1,234.50"],
    ], encoding=encoding, delimiter=delimiter)
    result = parser.parse_document(content, "fund.tsv" if delimiter == "\t" else "fund.csv", "tencent_licaitong", "trade")
    fields = result["rows"][0]["fields"]
    assert fields["action"] == expected and fields["asset_type"] == "fund"
    assert fields["business_date"] == "2026-10-27" and fields["confirmed_date"] == "2026-10-28"
    assert fields["gross_amount"] == "1234.50" and fields["quantity"] == "12.50"
    assert result["rows"][0]["source"]["row"] == 4


@pytest.mark.parametrize("direction,status", [("未知方向", "部分成交"), ("", ""), ("撤单", "失败")])
def test_unknown_directions_statuses_not_inferred_from_execution_headers(direction, status):
    result = parse([["成交日期", "成交金额", "股票代码", "方向", "状态"], ["20261028", "10", "000001", direction, status]])
    fields = result["rows"][0]["fields"]
    assert fields["action"] is None and fields["status"] is None
    assert ("action", "missing") in codes(result) and ("status", "missing") in codes(result)


def test_no_generic_confirmation_or_asset_class_guess_and_stock_subscription_is_ambiguous():
    result = parse([["日期", "证券代码", "方向"], ["2026-10-28", "000001", "买入"]])
    fields = result["rows"][0]["fields"]
    assert fields["status"] is None and fields["asset_type"] is None
    result = parse([["股票代码", "方向"], ["000001", "申购"]])
    assert result["rows"][0]["fields"]["action"] is None
    assert ("action", "ambiguous_action") in codes(result)


@pytest.mark.parametrize("header,basis,cost", [("成本金额", "unknown", "200"), ("报表成本", "unknown", "200"),
    ("持仓成本", "holding_cost", "200"), ("摊薄成本金额", "diluted_cost", "200"), ("摊薄成本价", "diluted_cost", None),
    ("成本价", None, None)])
def test_total_cost_unit_cost_and_cost_basis_remain_separate(header, basis, cost):
    result = parse([["基金代码", "持仓数量", header], ["000001", "20", "200"]], kind="holding")
    fields = result["rows"][0]["fields"]
    assert fields["cost_basis_type"] == basis and fields["cost"] == cost and fields["price"] is None
    if basis == "diluted_cost":
        assert ("cost_basis_type", "diluted_cost") in codes(result)
    if basis == "unknown":
        assert ("cost_basis_type", "unknown_cost_basis") in codes(result)


def test_available_quantity_requires_explicit_mapping_and_input_is_unchanged():
    rows = [["基金代码", "可用数量", "成本价"], ["000003", "12", "1.5"]]
    result = parse(rows, kind="holding")
    assert result["rows"][0]["fields"]["quantity"] is None
    assert ("quantity", "missing") in codes(result)
    assert result["warnings"][0]["code"] == "headers_changed"
    mapping = {"quantity": "可用数量"}
    original = deepcopy(mapping)
    content = csv_bytes(rows)
    before = content[:]
    result = parser.parse_document(content, "holding.csv", "tencent_licaitong", "holding", mapping)
    assert result["rows"][0]["fields"]["quantity"] == "12"
    assert result["rows"][0]["fields"]["cost"] is None
    assert mapping == original and content == before
    assert result == parser.parse_document(content, "holding.csv", "tencent_licaitong", "holding", mapping)


def test_missing_mapped_header_does_not_fall_back_and_ambiguous_headers_need_selection():
    rows = [["日期", "成交日期", "总份额", "股票代码", "方向"], ["20261027", "20261028", "10", "000001", "买入"]]
    assert ("business_date", "ambiguous_header") in {(w["field"], w["code"]) for w in parse(rows)["warnings"]}
    result = parse(rows, mapping={"business_date": "成交日期", "quantity": "旧份额", "status": "旧状态"})
    fields = result["rows"][0]["fields"]
    assert fields["business_date"] == "2026-10-28" and fields["quantity"] is None and fields["status"] is None
    assert any(w["code"] == "headers_changed" for w in result["warnings"])
    assert result["suggested_mapping"]["business_date"] == "成交日期"


@pytest.mark.parametrize("mapping", [{"bogus": "日期"}, {"price": 1}, {"price": ""}, {"price": None}, [], {"price": "金额", "cost": "金额"}])
def test_invalid_mappings_rejected(mapping):
    with pytest.raises(ValueError, match="Mapping|source header"):
        parse([["日期"], ["20261028"]], mapping=mapping)


@pytest.mark.parametrize("raw,expected", [("￥1,234.50", "1234.50"), ("(1,234.50)", "-1234.50"),
    ("(￥1,234.50)", "-1234.50"), ("￥(1,234.50)", "-1234.50"),
    ("-12.00", "-12.00"), ("0", "0"), ("", None), ("--", None), ("1e2", "100"), (".5", "0.5")])
def test_decimal_normalization_preserves_missing_and_explicit_zero(raw, expected):
    result = parse([["手续费"], [raw]])
    if raw == "":
        assert result["rows"] == []
    else:
        assert result["rows"][0]["fields"]["fee_amount"] == expected


@pytest.mark.parametrize("raw", ["NaN", "Infinity", "-inf", "1e999", "1,23", "$10", "1.2.3"])
def test_invalid_numbers_are_missing_with_issue(raw):
    result = parse([["发生金额"], [raw]])
    assert result["rows"][0]["fields"]["gross_amount"] is None
    assert ("gross_amount", "invalid_value") in codes(result)


def test_cash_aliases_negative_amount_missing_date_and_no_fees_computed():
    result = parse([["业务类型", "净收付", "币种", "对方账户", "流水号"], ["出金", "(1,000.00)", "人民币", "00001234", "cash-1"]], kind="cash")
    fields = result["rows"][0]["fields"]
    assert fields["business_date"] is None and fields["fee_amount"] is None and fields["gross_amount"] is None
    assert fields["asset_type"] == "cash" and fields["action"] == "withdraw" and fields["net_cash"] == "-1000.00"
    assert fields["counter_account_code"] == "00001234" and fields["source_ref"] == "cash-1"


@pytest.mark.parametrize("field,raw", [("日期", "2026-02-30"), ("成交时间", "25:00:00"), ("成交时间", "10:00:00+08:00"),
    ("日期", "2026-10-28 invalid"), ("证券代码", "１２３４５６"),
    ("证券代码", "6000000"), ("币种", "USD"), ("记录类型", "未知种类")])
def test_bad_dates_times_codes_and_enums_remain_missing(field, raw):
    result = parse([[field], [raw]])
    canonical = parser.HEADER_FIELDS[field]
    assert result["rows"][0]["fields"][canonical] is None
    assert (canonical, "invalid_value") in codes(result)


def test_xlsx_uses_safe_options_dates_sheet_row_and_numeric_code(monkeypatch):
    original = parser.load_workbook
    options = []

    def checked_load(*args, **kwargs):
        options.append(kwargs)
        return original(*args, **kwargs)

    monkeypatch.setattr(parser, "load_workbook", checked_load)
    content = xlsx_bytes([["报表"], ["成交日期", "成交时间", "股票代码", "方向", "成交金额"],
        [datetime(2026, 10, 28), time(10, 11, 12), 1, "卖出", 15.25]])
    row = parser.parse_document(content, "account.xlsx", "tonghuashun", "trade")["rows"][0]
    assert row["fields"]["business_date"] == "2026-10-28" and row["fields"]["trade_time"] == "10:11:12"
    assert row["fields"]["symbol"] == "000001" and row["fields"]["status"] == "confirmed"
    assert row["source"]["sheet"] == "Sheet1" and row["source"]["row"] == 3
    assert options == [{"read_only": True, "data_only": False, "keep_links": False}]


@pytest.mark.parametrize("rows", [[["股票代码", "成交金额"], ["000001", "=1+1"]], [["=1+1"], ["股票代码"], ["000001"]]])
def test_formula_anywhere_in_xlsx_rejected(rows):
    with pytest.raises(ValueError, match="Formula"):
        parser.parse_document(xlsx_bytes(rows), "account.xlsx", "tonghuashun", "trade")


@pytest.mark.parametrize("entries", [
    [("xl/vbaProject.bin", b"macro")],
    [("[Content_Types].xml", b'<Types ContentType="application/vnd.ms-excel.sheet.macroEnabled.main+xml"/>')],
    [("xl/externalLinks/externalLink1.xml", b"<externalLink/>")],
    [("xl/_rels/workbook.xml.rels", b'<Relationships><Relationship TargetMode="External" Target="https://example.invalid/"/></Relationships>')],
    [("xl/worksheets/sheet1.xml", b'<!DOCTYPE a [<!ENTITY x "123">]><a>&x;</a>')],
])
def test_macros_external_links_and_xml_entities_rejected_before_load(entries, monkeypatch):
    monkeypatch.setattr(parser, "load_workbook", lambda *a, **kw: pytest.fail("Unsafe archive reached openpyxl"))
    with pytest.raises(ValueError):
        parser.parse_document(zip_bytes(entries), "account.xlsx", "tonghuashun", "trade")


def test_zip_ratio_expanded_size_sparse_row_and_sheet_bounds():
    assert (parser.MAX_ROWS, parser.MAX_SHEETS, parser.MAX_CELLS) == (2000, 10, 100000)
    for content, message in [
        (zip_bytes([("huge.txt", b"0" * 100000)]), "ratio"),
        (zip_bytes([("huge.txt", b"0" * (40 * 1024**2 + 1))]), "expanded"),
        (zip_bytes([("xl/worksheets/sheet1.xml", b'<worksheet><row r="1048576"><c r="A1048576"/></row></worksheet>')]), "coordinate"),
        (xlsx_bytes([["股票代码"], ["000001"]], sheets=11), "Sheet"),
    ]:
        with pytest.raises(ValueError, match=message):
            parser.parse_document(content, "account.xlsx", "tonghuashun", "holding")


def test_row_limit_is_total_across_sheets_and_csv_accepts_boundary():
    rows = [["股票代码"], *[["000001"] for _ in range(2000)]]
    assert len(parse(rows, kind="holding")["rows"]) == 2000
    with pytest.raises(ValueError, match="row"):
        parse(rows + [["000002"]], kind="holding")
    content = xlsx_bytes([["股票代码"], *[["000001"] for _ in range(1001)]], sheets=2)
    with pytest.raises(ValueError, match="row"):
        parser.parse_document(content, "account.xlsx", "tonghuashun", "holding")


@pytest.mark.parametrize("extension", ["csv", "xlsx"])
def test_cell_limit_and_blank_padded_rows_are_bounded(extension, monkeypatch):
    monkeypatch.setattr(parser, "MAX_CELLS", 4)
    rows = [["股票代码", "数量"], ["000001", "1"]]
    encode = csv_bytes if extension == "csv" else xlsx_bytes
    assert len(parser.parse_document(encode(rows), f"a.{extension}", "tonghuashun", "holding")["rows"]) == 1
    with pytest.raises(ValueError, match="[Cc]ell|cell"):
        parser.parse_document(encode(rows + [["000002", "2"]]), f"a.{extension}", "tonghuashun", "holding")


@pytest.mark.parametrize("content,filename", [(b"not a zip", "a.xlsx"), (b'"unclosed', "a.csv"), (b"a\x00b", "a.csv"),
    (b'{"rows": []}', "a.csv"), (b"{}", "a.json"), (b"anything", "a.xlsm"), (b"", "a.csv"),
    (b"x" * (10 * 1024**2 + 1), "a.csv")], ids=["bad_zip", "bad_quotes", "nul", "disguised_json", "json", "macro_extension", "empty", "oversized"])
def test_malformed_unsupported_and_oversized_documents(content, filename):
    with pytest.raises(ValueError):
        parser.parse_document(content, filename, "tonghuashun", "trade")


@pytest.mark.parametrize("extension", ["PNG", "jpg", "jpeg"])
def test_images_only_report_ocr_pending(extension):
    result = parser.parse_document(b"image bytes", f"screenshot.{extension}", "tencent_licaitong", "holding")
    assert result["rows"] == [] and result["headers"] == [] and result["suggested_mapping"] == {}
    assert result["warnings"] == [{"field": None, "code": "ocr_pending", "message": "Image parsing awaits OCR; no values inferred", "severity": "warning"}]


def test_duplicate_headers_row_shape_and_empty_document():
    with pytest.raises(ValueError, match="Duplicate"):
        parse([["日期", "日期"], ["20261028", "20261028"]])
    result = parse([["股票代码", ""], ["000001", "unlabelled", "extra"]])
    assert (None, "row_shape") in codes(result)
    assert parser.parse_document(b"\r\n", "empty.csv", "tonghuashun", "trade")["rows"] == []


@pytest.mark.parametrize("profile,kind", [("other", "trade"), ("tonghuashun", "other")])
def test_unknown_profile_or_kind_rejected(profile, kind):
    with pytest.raises(ValueError, match="profile or kind"):
        parser.parse_document(b"x", "a.csv", profile, kind)


@pytest.mark.parametrize("header", ["税费", "税费合计", "金额税费total"])
def test_aggregate_charges_are_fee_not_tax_and_breakdown_unknown(header):
    result = parse([["股票代码", header], ["000001", "0.75"]])
    fields = result["rows"][0]["fields"]
    assert fields["fee_scope"] == "total_charges" and fields["fee_amount"] == "0.75"
    assert fields["tax_amount"] is None


def test_separate_fees_do_not_imply_zero_tax_and_total_scope_never_double_counts():
    fields = parse([["股票代码", "手续费"], ["000001", "0"]])["rows"][0]["fields"]
    assert fields["fee_scope"] == "separate" and fields["fee_amount"] == "0" and fields["tax_amount"] is None
    result = parse([["股票代码", "税费合计", "印花税", "费用口径"], ["000001", "10", "2", "separate"]])
    fields = result["rows"][0]["fields"]
    assert fields["fee_scope"] == "total_charges" and fields["fee_amount"] == "10" and fields["tax_amount"] is None
    assert ("fee_scope", "scope_conflict") in codes(result) and ("tax_amount", "aggregate_breakdown") in codes(result)
    result = parse([["股票代码", "费用", "费用口径", "印花税"], ["000001", "10", "total_charges", "2"]])
    assert result["rows"][0]["fields"]["tax_amount"] is None


def test_common_fields_and_snapshot_precision_are_preserved_for_parent_preflight():
    rows = [["基金代码", "总份额", "价格", "市值", "持仓成本", "asset_id", "note", "duplicate_decision"],
        ["000001", "1.123456789", "2.123456789", "3.123456", "4.123456", "0007", "人工复核", "keep"]]
    fields = parse(rows, kind="holding")["rows"][0]["fields"]
    assert {k: fields[k] for k in ("quantity", "price", "market_value", "cost")} == {
        "quantity": "1.123456789", "price": "2.123456789", "market_value": "3.123456", "cost": "4.123456"}
    assert fields["asset_id"] == "0007" and fields["note"] == "人工复核" and fields["duplicate_decision"] == "keep"


def test_exchange_fund_buy_is_not_otc_subscription_and_cost_columns_do_not_contaminate():
    rows = [["证券代码", "证券类型", "方向"], ["510300", "ETF", "买入"]]
    fields = parse(rows)["rows"][0]["fields"]
    assert fields["asset_type"] == "fund" and fields["action"] == "buy"
    result = parse([["股票代码", "成本金额", "摊薄成本价"], ["000001", "1000", "10"]], kind="holding")
    assert result["rows"][0]["fields"]["cost"] == "1000"
    assert result["rows"][0]["fields"]["cost_basis_type"] == "unknown"
    result = parse([["股票代码", "持仓成本", "成本类型"], ["000001", "1000", "未知"]], kind="holding")
    assert result["rows"][0]["fields"]["cost_basis_type"] == "unknown"


@pytest.mark.parametrize("reference", ["ZZZZZZZ1", "XFE1", "A0", "A9999999999"])
def test_sparse_cell_coordinates_rejected_before_openpyxl(reference, monkeypatch):
    monkeypatch.setattr(parser, "load_workbook", lambda *a, **kw: pytest.fail("Unbounded coordinates reached openpyxl"))
    content = zip_bytes([("xl/worksheets/sheet1.xml", f'<worksheet><row r="1"><c r="{reference}"/></row></worksheet>')])
    with pytest.raises(ValueError, match="coordinate"):
        parser.parse_document(content, "account.xlsx", "tonghuashun", "holding")


def test_utf16_xml_entities_and_csv_formula_rejected():
    xml = '<!DOCTYPE a [<!ENTITY x "123">]><a>&x;</a>'.encode("utf-16-le")
    with pytest.raises(ValueError, match="Unsafe"):
        parser.parse_document(zip_bytes([("xl/worksheets/sheet1.xml", xml)]), "a.xlsx", "tonghuashun", "holding")
    with pytest.raises(ValueError, match="Formula"):
        parse([["证券名称"], ["=WEBSERVICE(\"https://example.invalid/\")"]])


def test_multisheet_headers_changed_and_ambiguous_status_cannot_imply_confirmation():
    book = Workbook()
    book.active.append(["股票代码", "数量"])
    book.active.append(["000001", "10"])
    book.create_sheet("Funds").append(["基金代码", "总份额"])
    book["Funds"].append(["000002", "12"])
    content = BytesIO()
    book.save(content)
    book.close()
    result = parser.parse_document(content.getvalue(), "a.xlsx", "tonghuashun", "holding")
    assert len(result["rows"]) == 2 and any(w["code"] == "headers_changed" for w in result["warnings"])
    result = parse([["成交日期", "成交金额", "股票代码", "方向", "状态", "确认状态"],
        ["20261028", "1", "000001", "买入", "未知", "未知"]])
    assert result["rows"][0]["fields"]["status"] is None


def test_empty_execution_values_do_not_prove_confirmation_and_blank_padding_is_not_error():
    result = parse([["成交日期", "成交金额", "股票代码", "方向"], ["", "", "000001", "买入", ""]])
    assert result["rows"][0]["fields"]["status"] is None
    assert (None, "row_shape") not in codes(result)
    assert ("status", "missing") in codes(result)


def test_ten_sheets_are_allowed_and_formula_in_hidden_sheet_still_rejected():
    content = xlsx_bytes([["股票代码"], ["000001"]], sheets=10)
    assert len(parser.parse_document(content, "a.xlsx", "tonghuashun", "holding")["rows"]) == 10
    book = Workbook()
    book.active.append(["股票代码"])
    book.active.append(["000001"])
    hidden = book.create_sheet("Hidden")
    hidden.sheet_state = "hidden"
    hidden.append(["=1+1"])
    content = BytesIO()
    book.save(content)
    book.close()
    with pytest.raises(ValueError, match="Formula"):
        parser.parse_document(content.getvalue(), "a.xlsx", "tonghuashun", "holding")


def test_remapping_status_or_currency_to_note_cannot_enable_inferred_defaults():
    rows = [["成交日期", "股票代码", "方向", "成交金额", "状态"], ["20261028", "000001", "买入", "100", "部分成交"]]
    result = parse(rows, mapping={"note": "状态"})
    assert result["rows"][0]["fields"]["status"] is None
    assert ("status", "missing") in codes(result)
    fields = parse([["股票代码", "币种"], ["000001", "USD"]], mapping={"note": "币种"})["rows"][0]["fields"]
    assert fields["currency"] is None


def test_mixed_identifier_headers_need_asset_type_evidence_and_raw_text_keeps_spaces():
    rows = [["股票代码", "基金名称"], [" 000001 ", " 测试 "]]
    row = parse(rows)["rows"][0]
    assert row["fields"]["asset_type"] is None
    assert row["fields"]["symbol"] == "000001" and row["fields"]["name"] == "测试"
    assert row["source"]["raw_text"] == " 000001 \t 测试 "
